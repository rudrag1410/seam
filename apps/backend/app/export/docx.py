import io
from typing import cast

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.styles.style import _ParagraphStyle

from app.export.layout import to_blocks
from app.schemas.enums import BlockKind, ExportTemplate

_FONTS = {ExportTemplate.CLASSIC: "Times New Roman", ExportTemplate.MODERN: "Calibri"}


def render_docx(doc: dict, template: ExportTemplate = ExportTemplate.CLASSIC) -> bytes:
    d = Document()
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Pt(54)
        sec.top_margin = sec.bottom_margin = Pt(46)
    normal = cast(_ParagraphStyle, d.styles["Normal"])
    normal.font.name = _FONTS.get(template, "Calibri")
    normal.font.size = Pt(10.5)

    for b in to_blocks(doc):
        if b.kind is BlockKind.NAME:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(b.text).font.size = Pt(20)
            p.runs[0].bold = True
        elif b.kind is BlockKind.CONTACT:
            p = d.add_paragraph(b.text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b.kind is BlockKind.HEADING:
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run(b.text.upper())
            r.bold = True
            r.font.size = Pt(11)
        elif b.kind is BlockKind.ROLE:
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.add_run(b.text).bold = True
            if b.meta:
                p.add_run(f"   {b.meta}")
        elif b.kind is BlockKind.BULLET:
            d.add_paragraph(b.text, style="List Bullet")
        else:
            d.add_paragraph(b.text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
