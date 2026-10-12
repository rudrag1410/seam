import io

from docx import Document

from app.core.config import Settings
from app.core.errors import AppError
from app.parsing.pdf import ExtractedDocument

_WORDS_PER_PAGE = 500


def extract_docx(data: bytes, settings: Settings) -> ExtractedDocument:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise AppError(
            "FILE_CORRUPT", "We couldn't open that DOCX.", hint="Re-export it and try again."
        ) from exc

    lines: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        is_list = (p.style.name or "").lower().startswith("list") if p.style is not None else False
        lines.append(f"• {text}" if is_list else text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))

    text = "\n".join(lines)
    links = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.reltype.endswith("/hyperlink") and rel.target_ref.startswith(("http://", "https://"))
    ]
    pages = max(1, round(len(text.split()) / _WORDS_PER_PAGE))
    if pages > settings.max_pages:
        raise AppError(
            "TOO_MANY_PAGES",
            f"Resumes over {settings.max_pages} pages aren't supported. Is this the right file?",
        )
    return ExtractedDocument(text, pages, [text], links)
